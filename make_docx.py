#!/usr/bin/env python3
"""Generate editable .docx version of the Ledger Build vs Buy decision note."""
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY = "0F2A4A"; GREEN = "2E7D44"; RED = "B3392F"; ORANGE = "9A4E0E"
GREY_BG = "F1F5F9"; RED_BG = "FEF2F2"; BLUE_BG = "EAF2FB"; GREEN_BG = "F0FAF2"
DASH_BG = "FBFDFF"; GAP_BG = "FFF7F0"; DIM_BG = "EEF3F8"; ALT_BG = "F7FAFC"

doc = Document()
sec = doc.sections[0]
sec.page_height, sec.page_width = Mm(297), Mm(210)
sec.top_margin = sec.bottom_margin = Mm(12)
sec.left_margin = sec.right_margin = Mm(12)

style = doc.styles["Normal"]
style.font.name = "Calibri"; style.font.size = Pt(9)
style.paragraph_format.space_after = Pt(0)
style.paragraph_format.space_before = Pt(0)

def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)

def borders(cell, color="DBE3EC", sz=4, edges=("top","bottom","left","right")):
    tcPr = cell._tc.get_or_add_tcPr()
    tb = OxmlElement("w:tcBorders")
    for e in edges:
        el = OxmlElement(f"w:{e}")
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), str(sz))
        el.set(qn("w:color"), color)
        tb.append(el)
    tcPr.append(tb)

def no_borders(table):
    tbl = table._tbl
    pr = tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for e in ("top","left","bottom","right","insideH","insideV"):
        el = OxmlElement(f"w:{e}"); el.set(qn("w:val"), "none")
        b.append(el)
    pr.append(b)

def cell_text(cell, runs, align="center", size=8, space_before=2, space_after=2):
    """runs: list of (text, bold, colorhex) tuples; '\n' starts new paragraph."""
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = cell.paragraphs[0]
    p.alignment = {"center": WD_ALIGN_PARAGRAPH.CENTER, "left": WD_ALIGN_PARAGRAPH.LEFT}[align]
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    for text, bold, color in runs:
        parts = text.split("\n")
        for i, part in enumerate(parts):
            if i > 0:
                p = cell.add_paragraph()
                p.alignment = {"center": WD_ALIGN_PARAGRAPH.CENTER, "left": WD_ALIGN_PARAGRAPH.LEFT}[align]
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(space_after)
            r = p.add_run(part)
            r.bold = bold; r.font.size = Pt(size)
            r.font.color.rgb = RGBColor.from_string(color)

def heading(num, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7); p.paragraph_format.space_after = Pt(4)
    r = p.add_run(f" {num} ")
    r.bold = True; r.font.size = Pt(10); r.font.color.rgb = RGBColor.from_string("FFFFFF")
    r.font.highlight_color = None
    rPr = r._element.get_or_add_rPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), NAVY)
    rPr.append(shd)
    r2 = p.add_run("  " + text)
    r2.bold = True; r2.font.size = Pt(11); r2.font.color.rgb = RGBColor.from_string(NAVY)
    pb = OxmlElement("w:pBdr")
    btm = OxmlElement("w:bottom")
    btm.set(qn("w:val"), "single"); btm.set(qn("w:sz"), "8"); btm.set(qn("w:color"), "E2E8F0")
    pb.append(btm)
    p._p.get_or_add_pPr().append(pb)

def flow_row(steps, widths):
    """steps: list of (title, body, bg, border) or '->' arrows."""
    t = doc.add_table(rows=1, cols=len(steps))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    no_borders(t)
    for i, s in enumerate(steps):
        c = t.rows[0].cells[i]
        c.width = Mm(widths[i])
        if s == "->":
            cell_text(c, [("→", True, "64748B")], size=12)
        else:
            title, body, bg, bd = s
            shade(c, bg); borders(c, bd, sz=8)
            cell_text(c, [(title, True, "1E293B"), ("\n" + body, False, "334155")], size=7.5)
    return t

def strip(text_runs, bg):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    no_borders(t)
    c = t.rows[0].cells[0]; c.width = Mm(175)
    shade(c, bg)
    cell_text(c, text_runs, size=8.5, space_before=3, space_after=3)
    return t

def spacer(pts=4):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(pts)
    p.add_run("").font.size = Pt(2)

# ---------- TITLE BAR ----------
t = doc.add_table(rows=1, cols=1); no_borders(t)
c = t.rows[0].cells[0]; c.width = Mm(186); shade(c, NAVY)
cell_text(c, [
    ("Marketplace Ledger — Build vs Buy", True, "FFFFFF"),
], align="left", size=16, space_before=6, space_after=0)
p = c.add_paragraph(); p.paragraph_format.space_after = Pt(6)
r = p.add_run("Decision note for go-ahead   •   Recommendation: build the ledger in-house as the system of record   •   June 2026")
r.font.size = Pt(8.5); r.font.color.rgb = RGBColor.from_string("BCD3EC")

def labeled_arrow(cell, label, label_color, label_bg):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    shade(cell, "FFFFFF")
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(0)
    r = p.add_run(label)
    r.bold = True; r.font.size = Pt(6.5); r.font.color.rgb = RGBColor.from_string(label_color)
    rPr = r._element.get_or_add_rPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), label_bg)
    rPr.append(shd)
    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(2)
    r2 = p2.add_run("→")
    r2.bold = True; r2.font.size = Pt(12); r2.font.color.rgb = RGBColor.from_string("64748B")

# ---------- SECTION 1 ----------
heading("1", "How money flows today (current process)")

# Money relationship strip: AR on the LSP side, AP on the FO side
t = doc.add_table(rows=1, cols=5)
t.alignment = WD_TABLE_ALIGNMENT.CENTER
no_borders(t)
mwidths = [Mm(48), Mm(28), Mm(52), Mm(28), Mm(48)]
for j, w in enumerate(mwidths):
    t.rows[0].cells[j].width = w
c = t.rows[0].cells[0]
shade(c, BLUE_BG); borders(c, "6F9FD0", sz=8)
cell_text(c, [("LSP (customer)", True, "1E293B"),
              ("\nBooks the trip and owes us freight", False, "334155")], size=7.5)
labeled_arrow(t.rows[0].cells[1], " Money IN — Receivable (AR) ", "1F5C31", "E2F3E6")
c = t.rows[0].cells[2]
shade(c, GREY_BG); borders(c, "94A3B8", sz=8)
cell_text(c, [("Freight Tiger (marketplace)", True, "1E293B"),
              ("\nSits in the middle — collects from LSP, pays FO, keeps the margin", False, "334155")], size=7.5)
labeled_arrow(t.rows[0].cells[3], " Money OUT — Payable (AP) ", "7A3306", "FDEBD7")
c = t.rows[0].cells[4]
shade(c, BLUE_BG); borders(c, "6F9FD0", sz=8)
cell_text(c, [("Fleet Owner (vendor)", True, "1E293B"),
              ("\nRuns the trip and is owed the payout", False, "334155")], size=7.5)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(4)
r = p.add_run("Today both sides are mixed into the same trip fields — there is no separate “money in” (AR) vs “money out” (AP) view, so margin is invisible.")
r.font.size = Pt(7.5); r.italic = True; r.font.color.rgb = RGBColor.from_string("64748B")

flow_row([
    ("Trip created", "Ops team sets up the trip in CRM", GREY_BG, "94A3B8"), "->",
    ("Amounts typed into trip fields", "Freight, charges, deductions — each edit overwrites the old value", RED_BG, "EF9A9A"), "->",
    ("Payments recorded one-to-one", "One payment must match one charge; no part-payment or wallet", RED_BG, "EF9A9A"), "->",
    ("Finance cleans up manually", "Payments matched to trips in spreadsheets; Bill of Supply & TDS by hand", RED_BG, "EF9A9A"), "->",
    ("Reports cleaned manually", "Metabase data corrected before use", GREY_BG, "94A3B8"),
], [33, 5, 35, 5, 35, 5, 35, 5, 33])
spacer(2)
strip([("Result: ", True, "FFFFFF"),
       ("500+ finance hours every month on manual work  •  disputed balances we cannot defend  •  payments are our #1 support-ticket driver", False, "FFFFFF")], RED)

# ---------- SECTION 2 ----------
heading("2", "Gaps in the current process and system")
gaps = [
    ("No history.", " Edits overwrite old values, so we cannot show what a balance was on a given date. When a payment is disputed, we have nothing to point to."),
    ("Balances drift.", " Totals across a trip stop adding up over time, creating reconciliation mismatches between us, fleet owners and LSPs."),
    ("Rigid payments.", " One payment cannot settle many charges, be paid in parts, or sit in a wallet. Overpayments and cancelled-trip advances have no clean home."),
    ("No clear per-party picture.", " We can’t reliably see what each LSP owes us, what we owe each fleet owner, or our margin per trip."),
    ("Compliance is manual.", " Bill of Supply, credit notes and TDS are computed and issued by hand — ~37 compliance line items, all error-prone."),
    ("Heavy ongoing cost.", " All of the above lands on the finance team as repeated manual effort, month after month."),
]
t = doc.add_table(rows=3, cols=2)
t.alignment = WD_TABLE_ALIGNMENT.CENTER
no_borders(t)
for i, (head, body) in enumerate(gaps):
    c = t.rows[i // 2].cells[i % 2]
    c.width = Mm(92)
    shade(c, GAP_BG); borders(c, "F0CDAA", sz=6)
    cell_text(c, [(head, True, ORANGE), (body, False, "1E293B")], align="left", size=8.5, space_before=3, space_after=3)

# ---------- SECTION 3 ----------
heading("3", "What is needed to solve this")
flow_row([
    ("CRM stays the front end", "Teams keep working exactly where they do today", GREY_BG, "94A3B8"), "->",
    ("Ledger: the single source of truth", "Every change is saved as a new entry — nothing is ever overwritten", BLUE_BG, "6F9FD0"), "->",
    ("Clear balances", "Receivable / payable & margin per trip, per party, as of any date", GREEN_BG, "7FC08C"), "->",
    ("Flexible settlement", "One payment → many charges, part-payments, FO & LSP wallets", GREEN_BG, "7FC08C"),
], [36, 6, 44, 6, 40, 6, 40])
spacer(1)
flow_row([
    ("Documents auto-generated", "Bill of Supply, credit notes, TDS — amounts frozen when issued", GREEN_BG, "7FC08C"), "->",
    ("Payments connected", "Razorpay payouts and bank-receipt matching, driven by the ledger", GREEN_BG, "7FC08C"), "->",
    ("Optional: Zoho downstream", "If statutory books are needed, send summarised entries to Zoho — not the live system", DASH_BG, "9BB6D3"),
], [46, 6, 46, 6, 52])
spacer(2)
strip([("Result: ", True, "FFFFFF"),
       ("defensible balances on any date  •  automated compliance  •  500+ finance hours/month back  •  ledger visible to FOs & LSPs", False, "FFFFFF")], GREEN)

# ---------- PAGE BREAK ----------
doc.add_page_break()

# ---------- SECTION 4 ----------
heading("4", "Build vs Buy — the decision in one view")
rows = [
    ("Fits our business model",
     "✓ Designed for a broker: per-trip receivable (LSP) and payable (FO) with our margin in between.",
     "✗ Treats us as one company issuing invoices. No concept of two counterparties on one trip, or our margin."),
    ("Wallets & flexible settlement",
     "✓ Native wallets for FOs and LSPs; one payment can cover many charges or be part-paid.",
     "✗ No per-party wallets; settlement forced into Zoho’s invoice/bill structure."),
    ("Work we still do either way",
     "✓ One system to maintain. Business logic already lives in the marketplace.",
     "✗ Same work, plus more: onboard every FO & LSP into Zoho, keep records in sync, format data to Zoho’s rules."),
    ("When something breaks",
     "✓ We see and fix issues in our own system, directly.",
     "✗ Sync failures still land on our team — we carry the break-fix burden but debug across two systems."),
    ("Speed of change",
     "✓ New charge types, tax rules, wallet rules: we change and ship.",
     "✗ Constrained to Zoho’s object model; changes wait on what Zoho allows."),
    ("Reliability of live balances",
     "✓ Balances computed in our own system, in real time.",
     "✗ Live balances depend on Zoho’s sync reliability and API limits (~15–24k API calls/month at current volume)."),
    ("Cost & control",
     "✓ No per-seat licences; runs on existing infrastructure; full control of audit trail, data retention, INR/TDS handling.",
     "✗ Per-seat licensing that grows with the team; audit trail and retention live outside our control."),
    ("Where Zoho does help",
     "Statutory books, if needed: receive summarised entries from our ledger downstream.",
     "✓ Strong at standard statutory accounting — the one place it genuinely fits."),
]
t = doc.add_table(rows=len(rows) + 1, cols=3)
t.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr = t.rows[0]
for j, (txt, bg) in enumerate([("What matters", NAVY), ("Build in-house (recommended)", GREEN), ("Buy — Zoho Books as the live system", NAVY)]):
    c = hdr.cells[j]; shade(c, bg); borders(c)
    cell_text(c, [(txt, True, "FFFFFF")], align="left", size=8.5, space_before=2, space_after=2)
widths = [Mm(32), Mm(77), Mm(77)]
for j in range(3):
    hdr.cells[j].width = widths[j]
for i, (dim, build, buy) in enumerate(rows):
    row = t.rows[i + 1]
    c0, c1, c2 = row.cells
    for j, c in enumerate((c0, c1, c2)):
        c.width = widths[j]; borders(c)
    shade(c0, DIM_BG)
    cell_text(c0, [(dim, True, NAVY)], align="left", size=8)
    shade(c1, GREEN_BG)
    mark, rest = (build[0], build[1:]) if build[0] in "✓✗" else ("", build)
    col = GREEN if mark == "✓" else RED
    cell_text(c1, ([(mark + " ", True, col)] if mark else []) + [(rest.lstrip(), False, "1E293B")], align="left", size=8)
    if i % 2 == 1: shade(c2, ALT_BG)
    mark2, rest2 = (buy[0], buy[1:]) if buy[0] in "✓✗" else ("", buy)
    col2 = GREEN if mark2 == "✓" else RED
    cell_text(c2, ([(mark2 + " ", True, col2)] if mark2 else []) + [(rest2.lstrip(), False, "1E293B")], align="left", size=8)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(3); p.paragraph_format.space_after = Pt(6)
r = p.add_run("The key point: buying does not remove the work. Onboarding, mapping, data formatting and reconciliation stay with us in both options — Zoho only adds a runtime dependency.")
r.font.size = Pt(8); r.font.color.rgb = RGBColor.from_string("64748B"); r.italic = True

# ---------- RECOMMENDATION ----------
t = doc.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER
no_borders(t)
c = t.rows[0].cells[0]; c.width = Mm(186)
shade(c, GREEN_BG); borders(c, GREEN, sz=10)
cell_text(c, [("Recommendation & ask", True, "1F5C31")], align="left", size=10, space_before=4)
p = c.add_paragraph(); p.paragraph_format.space_after = Pt(3)
r = p.add_run("Build the ledger natively in the marketplace as the system of record. ")
r.bold = True; r.font.size = Pt(8.7)
r = p.add_run("Every financial change is stored as a new entry (never overwritten), giving us defensible point-in-time balances, flexible settlement with wallets, a clean receivable/payable and margin view per party, and automated Bill of Supply, credit notes and TDS. If statutory books are required, we sync summarised entries to Zoho downstream rather than running it as the live system.")
r.font.size = Pt(8.7)
p = c.add_paragraph(); p.paragraph_format.space_after = Pt(4)
r = p.add_run("Ask: ")
r.bold = True; r.font.size = Pt(8.7)
r = p.add_run("go-ahead to start Phase 1 (ledger foundation, run in parallel with current process — zero business disruption), followed by compliance automation, reconciliation & reporting, and finally surfacing the ledger to fleet owners and LSPs.")
r.font.size = Pt(8.7)

doc.save("Ledger_Build_vs_Buy.docx")
print("saved")
