from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn

NAVY=RGBColor(0x1F,0x37,0x63); GREEN=RGBColor(0x26,0x81,0x5C)
PURPLE=RGBColor(0x60,0x43,0xA3); AMBER=RGBColor(0xC7,0x7D,0x0A)
SLATE=RGBColor(0x3E,0x4C,0x63); INK=RGBColor(0x25,0x30,0x3B)
GREY=RGBColor(0x6B,0x74,0x80); WHITE=RGBColor(0xFF,0xFF,0xFF)
BAND=RGBColor(0xF4,0xF6,0xF9); CARD=RGBColor(0xF7,0xF8,0xFA)
LINEC=RGBColor(0xD5,0xDA,0xE1)
FONT="Manrope"

prs=Presentation("base.pptx")

# pick a blank-ish layout
blank=min(prs.slide_layouts, key=lambda L: len(L.placeholders))

def set_font(run,size,bold,color,italic=False):
    run.font.size=Pt(size); run.font.bold=bold; run.font.italic=italic
    run.font.color.rgb=color; run.font.name=FONT
    rPr=run._r.get_or_add_rPr()
    for tag in ("a:latin","a:cs"):
        e=rPr.find(qn(tag))
        if e is None:
            e=rPr.makeelement(qn(tag),{}); rPr.append(e)
        e.set("typeface",FONT)

def add_slide():
    s=prs.slides.add_slide(blank)
    for ph in list(s.placeholders):
        ph._element.getparent().remove(ph._element)
    return s

def textbox(s,l,t,w,h,anchor=MSO_ANCHOR.TOP):
    tb=s.shapes.add_textbox(Emu(int(l)),Emu(int(t)),Emu(int(w)),Emu(int(h)))
    tf=tb.text_frame; tf.word_wrap=True; tf.vertical_anchor=anchor
    for m in ("margin_left","margin_right","margin_top","margin_bottom"):
        setattr(tf,m,Emu(0))
    return tb,tf

def para(tf,runs,first=False,align=PP_ALIGN.LEFT,space_before=0,space_after=2,
         line=1.0,bullet=False,indent=None):
    p=tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment=align; p.space_before=Pt(space_before); p.space_after=Pt(space_after)
    try: p.line_spacing=line
    except: pass
    for (text,size,bold,color,italic) in runs:
        r=p.add_run(); r.text=text; set_font(r,size,bold,color,italic)
    if bullet:
        pPr=p._p.get_or_add_pPr(); mar=indent or 130000
        pPr.set("indent",str(-mar)); pPr.set("marL",str(mar))
    return p

def rect(s,l,t,w,h,fill,rounded=False,line=None,lw=None,radius=0.09):
    shp=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                           Emu(int(l)),Emu(int(t)),Emu(int(w)),Emu(int(h)))
    if fill is None: shp.fill.background()
    else: shp.fill.solid(); shp.fill.fore_color.rgb=fill
    if line is None: shp.line.fill.background()
    else: shp.line.color.rgb=line; shp.line.width=Pt(lw or 0.75)
    shp.shadow.inherit=False
    if rounded:
        try: shp.adjustments[0]=radius
        except: pass
    return shp

ML=228600
def header(s,title,subtitle):
    _,tf=textbox(s,ML,150000,7900000,300000)
    para(tf,[(title,15,True,NAVY,False)],first=True,space_after=1)
    _,tf=textbox(s,ML,458000,8680000,220000)
    para(tf,[(subtitle,9,False,GREY,True)],first=True)
    # divider
    ln=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Emu(242398),Emu(662000),
                              Emu(242398+8616600),Emu(662000))
    ln.line.color.rgb=LINEC; ln.line.width=Pt(1.0); ln.shadow.inherit=False
    # logo
    s.shapes.add_picture("logo.png",Emu(8247328),Emu(65382),Emu(821475),Emu(355725))
    # footer
    _,tf=textbox(s,Emu(7182075),Emu(4905000),Emu(1652700),Emu(180000))
    para(tf,[("@FreightTiger confidential",8,False,RGBColor(0xB8,0xBE,0xC6),False)],
         first=True,align=PP_ALIGN.RIGHT)

# ============================================================
# SLIDE 3 — Key Challenges
# ============================================================
s=add_slide()
header(s,"Key Challenges — Scaling the Initiatives",
       "System-generated supply is actioned last — CRM inventory, FO App bids & bot demands need priority & SLA decisions from leadership")

c1=1750000; c2=3860000; c3=8686800-c1-c2
x0=ML; x1=x0+c1; x2=x1+c2
cols=[(x0,c1),(x1,c2),(x2,c3)]
hy=706000; hh=286000
hdr=[("CHALLENGE",NAVY),("CONTEXT",SLATE),("KEY QUESTIONS FOR LEADERSHIP",PURPLE)]
for (l,w),(txt,clr) in zip(cols,hdr):
    rect(s,l,hy,w-24000,hh,clr,rounded=True)
    _,tf=textbox(s,l+46000,hy,w-70000,hh,anchor=MSO_ANCHOR.MIDDLE)
    para(tf,[(txt,9.5,True,WHITE,False)],first=True,space_after=0)

challenges=[
 dict(n="1", short="CRM inventory prioritisation",
      q="How should CRM-generated inventory be prioritised?",
      ctx="~1,300 CRM inventories/month (incl. ~800 Power-Lane), but conversions stay under 50/month. The fulfilment team clears its own + central-team inventory first; by the time it reaches system-generated CRM inventory, the load is usually already placed externally (CRM inventory is collected in the morning).",
      qs=["Should CRM inventory get equal priority in the workflow, not the last source?",
          "What SLA / process ensures it is actioned before the demand is fulfilled elsewhere?"]),
 dict(n="2", short="FO App bids not acted upon",
      q="Why are FO App bids not being acted upon?",
      ctx="We receive 30–55 FO App bids/day, but only 1–2 are converted — and 0 conversions in the last week. As with CRM inventory, the team focuses on manual + central inventory first; by the time app bids are reviewed, the demands have often already been fulfilled.",
      qs=["Should FO App bids have a defined response SLA / priority, so they are reviewed before demand is fulfilled through other channels?"]),
 dict(n="3", short="Bot-created demand actionability",
      q="Why is there low actionability on bot-created demands?",
      ctx="On power lanes, manually-created demands convert at 18% vs only 3% for bot-created demands. Bot-created demand is clearly not getting priority, while the team continues to focus on manually-created demands.",
      qs=["Should bot-created demands sit in the primary fulfilment queue, not a secondary source?",
          "What business-process changes ensure bot demands are acted on before they go stale?"]),
]
ry=1040000; rh=1225000; gap=22000
for i,c in enumerate(challenges):
    top=ry+i*(rh+gap)
    rect(s,ML,top,x2+c3-ML,rh,BAND if i%2==0 else WHITE,rounded=True)
    # challenge
    _,tf=textbox(s,x0+60000,top+40000,c1-100000,rh-70000,anchor=MSO_ANCHOR.MIDDLE)
    para(tf,[(c["n"]+"   ",16,True,NAVY,False),(c["short"],10,True,NAVY,False)],
         first=True,space_after=4,line=1.02)
    para(tf,[(c["q"],8,False,GREY,True)],space_after=0,line=1.02)
    # context
    _,tf=textbox(s,x1+64000,top+34000,c2-120000,rh-60000,anchor=MSO_ANCHOR.MIDDLE)
    para(tf,[(c["ctx"],8.4,False,INK,False)],first=True,space_after=0,line=1.08)
    # questions
    _,tf=textbox(s,x2+64000,top+34000,c3-110000,rh-60000,anchor=MSO_ANCHOR.MIDDLE)
    for j,qq in enumerate(c["qs"]):
        para(tf,[("›  ",8.6,True,PURPLE,False),(qq,8.4,False,INK,False)],
             first=(j==0),space_after=6,line=1.06,bullet=True,indent=150000)

# ============================================================
# SLIDE 4 — 3-Month Execution Plan / Roadmap
# ============================================================
s=add_slide()
header(s,"3-Month Execution Plan — Carrier Matching",
       "Scale revenue to ₹1.2 Cr/month by Sept · cut cost base 50%    |    Q2 exit: DAU → 500 · conversion → 20% · MS+FinOps −30%   (25% / 50% = FY27)")

# legend
BUCKETS=[("Ledger",NAVY),("Fulfilment %",GREEN),("DAU",PURPLE),("Reduce Ops Bandwidth",AMBER)]
lx=ML; ly=712000
for name,clr in BUCKETS:
    rect(s,lx,ly+14000,150000,150000,clr,rounded=True,radius=0.3)
    _,tf=textbox(s,lx+185000,ly,1900000,190000,anchor=MSO_ANCHOR.MIDDLE)
    w=len(name)
    para(tf,[(name,8.6,True,clr,False)],first=True,space_after=0)
    lx+= 185000 + int(w*57000) + 150000
# conversion trend note (right)
_,tf=textbox(s,Emu(5250000),Emu(ly),Emu(3665400),Emu(190000),anchor=MSO_ANCHOR.MIDDLE)
para(tf,[("Network-lane conversion:  ",8.2,True,SLATE,False),
         ("Apr 13.7% → May 11.2% → Jun 17.1%  →  Q2 exit 20%",8.2,False,INK,False)],
     first=True,align=PP_ALIGN.RIGHT,space_after=0)

# Month -> Goal -> rolled-up task summary  (goal, colour, #stories, summary)
def g(name,clr,n,summ): return (name,clr,n,summ)
JULY=[
 g("Ledger",NAVY,5,"Stand up the trip-level payment ledger — immutable Net Payable/Receivable record, wallet-based overpayment settlement, and automated TDS (capture, calculate & reverse)."),
 g("Fulfilment %",GREEN,3,"Act on demand faster — lane-based agent auto-assignment, Demand-Bot context building, and near-match inventory with match-quality on the demand list."),
 g("DAU",PURPLE,1,"Let FOs reach the PSA directly for placement from the app (PSA number in-app)."),
 g("Reduce Ops Bandwidth",AMBER,3,"Automate ops — advance payments for Gold/Silver FOs, pre-transit SIM-consent collection, and transit-delay detection & escalation."),
]
AUG=[
 g("Ledger",NAVY,4,"Take Ledger Phase-1 to production and bring finance into the FO App — 3-yr trip statement, trip-level TDS rate, and self-serve TDS declaration."),
 g("Fulfilment %",GREEN,5,"Prioritise placeable demand — auto CRM alerts, Demand-Actionability-Score sorting, vehicle heat-map, auto tracking-master updates, and in-CRM calling with callbacks."),
 g("DAU",PURPLE,1,"Extend ticketing to FOs not on the app via WhatsApp / IVR / agent."),
 g("Reduce Ops Bandwidth",AMBER,1,"Auto-share the loading location to drivers over call + WhatsApp."),
]
SEP=[
 g("Ledger",NAVY,1,"Add configurable LSP exposure limits with auto-disable on breach."),
 g("Fulfilment %",GREEN,2,"Scale fulfilment — agentic-AI calling for agents and a live LSP / PSA-CSM / city target dashboard."),
 g("DAU",PURPLE,1,"Capture trip charges directly in the FO App with an approval flow."),
 g("Reduce Ops Bandwidth",AMBER,2,"Automate compliance-document generation and POD audit via OCR."),
]
MONTHS=[("JULY",JULY,12),("AUGUST",AUG,11),("SEPTEMBER",SEP,6)]

colw=(8686800-2*46000)//3
top_hdr=980000; hh=250000
body_top=1290000; body_bot=4300000
slot=(body_bot-body_top)//4
for ci,(mn,goals,total) in enumerate(MONTHS):
    cl=ML+ci*(colw+46000)
    # column card
    rect(s,cl,top_hdr,colw,body_bot-top_hdr,CARD,rounded=True,line=LINEC,lw=0.5)
    # month header
    rect(s,cl,top_hdr,colw,hh,SLATE,rounded=True)
    _,tf=textbox(s,cl+80000,top_hdr,colw-100000,hh,anchor=MSO_ANCHOR.MIDDLE)
    para(tf,[(mn+"    ",11,True,WHITE,False),
             ("%d stories"%total,8,True,RGBColor(0xC5,0xCC,0xD6),False)],
         first=True,space_after=0)
    # goal blocks
    for gi,(gname,gc,n,summ) in enumerate(goals):
        gy=top_hdr+hh+gi*slot
        if gi>0:
            dv=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Emu(int(cl+70000)),Emu(int(gy)),
                                      Emu(int(cl+colw-70000)),Emu(int(gy)))
            dv.line.color.rgb=RGBColor(0xE4,0xE8,0xED); dv.line.width=Pt(0.5); dv.shadow.inherit=False
        rect(s,cl+56000,gy+42000,44000,slot-92000,gc,rounded=True,radius=0.4)  # accent bar
        _,tf=textbox(s,cl+140000,gy+34000,colw-200000,slot-58000,anchor=MSO_ANCHOR.MIDDLE)
        para(tf,[(gname+"   ",9,True,gc,False),
                 ("· %d %s"%(n,"story" if n==1 else "stories"),7,True,GREY,False)],
             first=True,space_after=3,line=1.0)
        para(tf,[(summ,7.6,False,INK,False)],space_after=0,line=1.08)

# notes-for-leadership band
nby=body_bot+40000
rect(s,ML,nby,8686800,530000,BAND,rounded=True)
_,tf=textbox(s,ML+90000,nby+34000,8686800-180000,530000-60000,anchor=MSO_ANCHOR.MIDDLE)
para(tf,[("Notes for leadership   ",8.6,True,NAVY,False),
         ("Q2 exit targets: DAU → 500 · network-lane conversion → 20% · MS+FinOps bandwidth −30%  (25% / 50% are FY27 annual targets).",
          8.2,False,INK,False)],first=True,space_after=3,line=1.05)
para(tf,[("Network-lane conversion (Q1 FY27 actuals): Apr 13.7% → May 11.2% → Jun 17.1% — entering Q2 at ~17%, the plan above targets a 20% exit.",
          8.2,False,INK,False)],space_after=0,line=1.05)

prs.save("deck_v7.pptx")
print("saved deck_v7.pptx  slides:",len(prs.slides._sldIdLst))
