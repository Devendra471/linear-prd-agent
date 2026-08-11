from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn
from copy import deepcopy

NAVY=RGBColor(0x1F,0x37,0x63); GREEN=RGBColor(0x26,0x81,0x5C)
PURPLE=RGBColor(0x60,0x43,0xA3); AMBER=RGBColor(0xC7,0x7D,0x0A)
MAROON=RGBColor(0x9C,0x2B,0x2B); TEAL=RGBColor(0x0F,0x76,0x6E); BLUE=RGBColor(0x1D,0x6F,0xA5)
SLATE=RGBColor(0x3E,0x4C,0x63); INK=RGBColor(0x25,0x30,0x3B)
GREY=RGBColor(0x6B,0x74,0x80); WHITE=RGBColor(0xFF,0xFF,0xFF)
BAND=RGBColor(0xF4,0xF6,0xF9); CARD=RGBColor(0xF7,0xF8,0xFA); LINEC=RGBColor(0xD5,0xDA,0xE1)
FONT="Manrope"
ML=228600

# ---------------- part 1: edit slide 1 headers + next-focus ----------------
# NB: always rebuild from the single-slide source (july_out.pptx) to avoid
# duplicating slide 2 if this script is re-run.
p=Presentation("july_out.pptx")
s1=p.slides[0]

def para_set_text(p_elem, text):
    runs=p_elem.findall(qn('a:r'))
    first=runs[0]
    t=first.find(qn('a:t'))
    if t is None:
        t=first.makeelement(qn('a:t'),{}); first.append(t)
    t.text=text
    for r in runs[1:]:
        p_elem.remove(r)

def rewrite(shape_id, bullets):
    for sh in s1.shapes:
        if sh.shape_id==shape_id:
            tf=sh.text_frame
            paras=[pp._p for pp in tf.paragraphs]
            template=deepcopy(paras[0])
            parent=paras[0].getparent()
            for pe in paras: parent.remove(pe)
            for b in bullets:
                clone=deepcopy(template)
                para_set_text(clone, b)
                parent.append(clone)
            return
    raise SystemExit("shape %s not found"%shape_id)

def set_header(shape_id, txt):
    for sh in s1.shapes:
        if sh.shape_id==shape_id:
            para_set_text(sh.text_frame.paragraphs[0]._p, txt); return

set_header(64, "PROGRESS  (thru Jul'26)")
set_header(66, "NEXT FOCUS  (Aug–Sept)")

rewrite(70,[   # FO App
 "›  Trip Statement — 3-yr ledger tab in FO App",
 "›  TDS rate visibility + self-serve declaration",
 "›  Ticketing for FOs not on app (WhatsApp/IVR)",
])
rewrite(74,[   # Fulfilment
 "›  Auto CRM alerts — bids / matches / approvals",
 "›  Sort demand by Demand Actionability Score",
 "›  Agentic AI calling for fulfilment agents",
])
# 79 (Inventory) left unchanged — not covered by new roadmap table
rewrite(83,[   # Cost to serve / Ops
 "›  Runner App — POD collection + OCR audit",
 "›  LSP App — self-serve trip/tracking/POD + tickets",
 "›  Automate compliance document generation",
])
rewrite(88,[   # Ledger
 "›  Auto-calc & reverse TDS; advance-pay auto-release (60%)",
 "›  Payment fraud checks + LSP exposure limits",
 "›  Self-serve balance payment for FOs",
])

# ---------------- part 2: new slide — Aug–Sept roadmap detail ----------------
blank=min(p.slide_layouts, key=lambda L: len(L.placeholders))
s2=p.slides.add_slide(blank)
for ph in list(s2.placeholders):
    ph._element.getparent().remove(ph._element)

def set_font(run,size,bold,color,italic=False):
    run.font.size=Pt(size); run.font.bold=bold; run.font.italic=italic
    run.font.color.rgb=color; run.font.name=FONT
    rPr=run._r.get_or_add_rPr()
    for tag in ("a:latin","a:cs"):
        e=rPr.find(qn(tag))
        if e is None: e=rPr.makeelement(qn(tag),{}); rPr.append(e)
        e.set("typeface",FONT)

def textbox(sl,l,t,w,h,anchor=MSO_ANCHOR.TOP):
    tb=sl.shapes.add_textbox(Emu(int(l)),Emu(int(t)),Emu(int(w)),Emu(int(h)))
    tf=tb.text_frame; tf.word_wrap=True; tf.vertical_anchor=anchor
    for m in ("margin_left","margin_right","margin_top","margin_bottom"):
        setattr(tf,m,Emu(0))
    return tb,tf

def para(tf,runs,first=False,align=PP_ALIGN.LEFT,space_before=0,space_after=2,
         line=1.0,bullet=False,indent=None):
    pa=tf.paragraphs[0] if first else tf.add_paragraph()
    pa.alignment=align; pa.space_before=Pt(space_before); pa.space_after=Pt(space_after)
    try: pa.line_spacing=line
    except: pass
    for (text,size,bold,color,italic) in runs:
        r=pa.add_run(); r.text=text; set_font(r,size,bold,color,italic)
    if bullet:
        pPr=pa._p.get_or_add_pPr(); mar=indent or 120000
        pPr.set("indent",str(-mar)); pPr.set("marL",str(mar))
    return pa

def rect(sl,l,t,w,h,fill,rounded=False,line=None,lw=None,radius=0.09):
    shp=sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                            Emu(int(l)),Emu(int(t)),Emu(int(w)),Emu(int(h)))
    if fill is None: shp.fill.background()
    else: shp.fill.solid(); shp.fill.fore_color.rgb=fill
    if line is None: shp.line.fill.background()
    else: shp.line.color.rgb=line; shp.line.width=Pt(lw or 0.75)
    shp.shadow.inherit=False
    if rounded:
        try: shp.adjustments[0]=radius
        except: pass
    return shp

def header(sl,title,subtitle):
    _,tf=textbox(sl,ML,155000,7900000,300000)
    para(tf,[(title,15,True,NAVY,False)],first=True,space_after=1)
    _,tf=textbox(sl,ML,458000,8680000,220000)
    para(tf,[(subtitle,9,False,GREY,True)],first=True)
    ln=sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Emu(242398),Emu(662000),
                               Emu(242398+8616600),Emu(662000))
    ln.line.color.rgb=LINEC; ln.line.width=Pt(1.0); ln.shadow.inherit=False
    sl.shapes.add_picture("logo.png",Emu(8247328),Emu(65382),Emu(821475),Emu(355725))
    _,tf=textbox(sl,Emu(7182075),Emu(4905000),Emu(1652700),Emu(180000))
    para(tf,[("@FreightTiger confidential",8,False,RGBColor(0xB8,0xBE,0xC6),False)],
         first=True,align=PP_ALIGN.RIGHT)

header(s2,"Next Focus — Aug–Sept'26 Roadmap",
       "15 initiatives across Finance/Ledger, Risk, Fulfilment, FO/LSP Experience & Ops — bucketed with measurable outcomes")

# bucket data: (bucket_name, color, [(task, outcome)])
BUCKETS = {
 "finance": ("Finance Automation & Ledger", NAVY, [
   ("Ledger Automation in CRM", "100% trips carry immutable payment record; overpayment auto-settled to wallets"),
   ("Auto-calculate TDS", "100% trips system-calculated; ~0 error rate; 60 hrs/mo ops saved"),
   ("Automated Advance Payment (Gold/Silver FOs)", "60% auto-released; disbursal TAT → 2 hrs"),
   ("Automate compliance document generation", "100% auto-generated; 60 hrs/mo saved"),
   ("Balance payment automation", "60% self-initiated by FO; fewer ops touchpoints"),
 ]),
 "fulfil": ("Fulfilment & Demand Intelligence", GREEN, [
   ("Auto CRM alerts", "bids received / inventory matched / approval assigned — faster time-to-action"),
   ("Sort demand by Demand Actionability Score", "more agent time on placeable loads; faster placement"),
   ("Agentic AI calling for fulfilment agents", "calls automated at scale; cost/call ↓; conversion ↑"),
 ]),
 "foapp": ("FO App Transparency & Self-Serve", PURPLE, [
   ("Trip Statement — 3-yr ledger tab in FO App", "statement views ↑; fewer payment-status tickets"),
   ("TDS rate visibility + self-serve declaration", "% FOs self-serving TDS ↑; fewer TDS tickets"),
 ]),
 "risk": ("Risk, Fraud & Exposure Control", MAROON, [
   ("Payment fraud checks (FO / user / bank-detail change)", "100% payments fraud-screened; releases blocked pre-disbursal"),
   ("Configurable LSP exposure limit + auto-disable", "breaches auto-caught & actioned; lower bad-debt risk"),
 ]),
 "pod": ("POD Collection & Audit", TEAL, [
   ("Runner App/Mobile page", "Hard-POD TAT ↓ ≥20%; OCR-audited PODs; faster dispute resolution"),
 ]),
 "lsp": ("LSP Experience & Visibility", BLUE, [
   ("LSP App", "self-serve trip/tracking/POD status + ticketing; fewer inbound status calls"),
 ]),
 "ops": ("Support & Ops Automation", AMBER, [
   ("Ticketing for FOs not on the app", "WhatsApp/IVR/agent — broader off-app support coverage"),
 ]),
}

col_w=(8686800-2*46000)//3
COLS=[
 [BUCKETS["finance"]],
 [BUCKETS["fulfil"], BUCKETS["foapp"]],
 [BUCKETS["risk"], BUCKETS["pod"], BUCKETS["lsp"], BUCKETS["ops"]],
]
top0=980000; bottom=4300000
gap=60000

for ci,cards in enumerate(COLS):
    cl=ML+ci*(col_w+46000)
    n_items=sum(len(c[2]) for c in cards)
    avail=bottom-top0-gap*(len(cards)-1)
    y=top0
    for name,color,items in cards:
        card_h = int(avail * (len(items)/n_items))
        rect(s2,cl,y,col_w,card_h,CARD,rounded=True,line=LINEC,lw=0.5)
        rect(s2,cl+18000,y+22000,36000,card_h-44000,color,rounded=True,radius=0.4)
        _,tf=textbox(s2,cl+80000,y+18000,col_w-120000,card_h-32000,anchor=MSO_ANCHOR.TOP)
        para(tf,[(name,9.5,True,color,False)],first=True,space_after=5,line=1.0)
        for j,(task,outcome) in enumerate(items):
            para(tf,[(task+"  ",7.6,True,INK,False)],space_after=1,line=1.05)
            para(tf,[(outcome,7.2,False,GREY,True)],space_after=6,line=1.05)
        y += card_h + gap

p.save("TML_July_Update.pptx")
print("saved TML_July_Update.pptx with slide2 roadmap; slides:",len(p.slides._sldIdLst))
