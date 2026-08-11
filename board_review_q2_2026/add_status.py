from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

SHIP=RGBColor(0x1E,0x7A,0x46); PART=RGBColor(0x0F,0x76,0x6E)
BUILD=RGBColor(0x3E,0x4C,0x63); PLAN=RGBColor(0x8A,0x93,0x9E)
FONT="Manrope"
STATUS={
 ("JULY","Ledger"):("✓ Shipped",SHIP),
 ("JULY","Fulfilment"):("◐ Part shipped",PART),
 ("JULY","DAU"):("◐ In build",BUILD),
 ("JULY","Reduce"):("◐ Part shipped",PART),
 ("AUGUST","Ledger"):("◐ In build",BUILD),
 ("AUGUST","Fulfilment"):("○ Planned",PLAN),
 ("AUGUST","DAU"):("○ Planned",PLAN),
 ("AUGUST","Reduce"):("○ Planned",PLAN),
 ("SEPTEMBER","Ledger"):("○ Planned",PLAN),
 ("SEPTEMBER","Fulfilment"):("○ Planned",PLAN),
 ("SEPTEMBER","DAU"):("○ Planned",PLAN),
 ("SEPTEMBER","Reduce"):("○ Planned",PLAN),
}
def month_of(left):
    return "JULY" if left<2_000_000 else ("AUGUST" if left<5_000_000 else "SEPTEMBER")
def goalkey(txt):
    for k in ("Ledger","Fulfilment","DAU","Reduce"):
        if txt.strip().startswith(k): return k
    return None
def add_run(para,text,size,bold,color):
    r=para.add_run(); r.text=text
    r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=color; r.font.name=FONT
    rPr=r._r.get_or_add_rPr()
    for tag in ("a:latin","a:cs"):
        e=rPr.find(qn(tag))
        if e is None: e=rPr.makeelement(qn(tag),{}); rPr.append(e)
        e.set("typeface",FONT)

p=Presentation("board_deck.pptx")
s=p.slides[3]  # roadmap
n=0
for sh in s.shapes:
    if not sh.has_text_frame: continue
    paras=sh.text_frame.paragraphs
    if not paras: continue
    p0=paras[0]; full="".join(r.text for r in p0.runs)
    if ("stories" in full or "story" in full) and goalkey(full):
        gk=goalkey(full); mn=month_of(sh.left or 0)
        stat,clr=STATUS.get((mn,gk),("○ Planned",PLAN))
        add_run(p0,"   "+stat,7,True,clr)
        n+=1

# status key on the notes band
for sh in s.shapes:
    if sh.has_text_frame and "Notes for leadership" in sh.text_frame.text:
        para=sh.text_frame.add_paragraph()
        para.space_before=Pt(3)
        add_run(para,"Status per Linear (Aug'26):  ",7.5,True,RGBColor(0x25,0x30,0x3B))
        add_run(para,"✓ Shipped   ",7.5,True,SHIP)
        add_run(para,"◐ In build / part shipped   ",7.5,True,PART)
        add_run(para,"○ Planned",7.5,True,PLAN)
        break

p.save("board_out.pptx")
print("added status to",n,"goal cells")
