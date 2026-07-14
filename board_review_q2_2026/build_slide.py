from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from PIL import Image
import copy

NAVY=RGBColor(0x1F,0x37,0x63); GREEN=RGBColor(0x26,0x81,0x5C)
PURPLE=RGBColor(0x60,0x43,0xA3); SLATE=RGBColor(0x3E,0x4C,0x63)
INK=RGBColor(0x25,0x30,0x3B); GREY=RGBColor(0x6B,0x74,0x80)
WHITE=RGBColor(0xFF,0xFF,0xFF)
BAND=RGBColor(0xF4,0xF6,0xF9); LINE=RGBColor(0xDD,0xE1,0xE7)
FONT="Manrope"

p=Presentation("deck.pptx")
s=p.slides[0]

# ---- keep logo, divider, page-number, footer; drop old table/title + edit-note(152) ----
KEEP={142,143,150,151}
for sh in list(s.shapes):
    if sh.shape_id not in KEEP:
        sh._element.getparent().remove(sh._element)

def set_font(run, size, bold, color, italic=False):
    run.font.size=Pt(size); run.font.bold=bold; run.font.italic=italic
    run.font.color.rgb=color; run.font.name=FONT
    rPr=run._r.get_or_add_rPr()
    for tag in ("a:latin","a:cs"):
        e=rPr.find(qn(tag))
        if e is None:
            e=rPr.makeelement(qn(tag),{}); rPr.append(e)
        e.set("typeface",FONT)

def textbox(l,t,w,h,anchor=MSO_ANCHOR.TOP):
    tb=s.shapes.add_textbox(Emu(int(l)),Emu(int(t)),Emu(int(w)),Emu(int(h)))
    tf=tb.text_frame; tf.word_wrap=True
    tf.vertical_anchor=anchor
    for m in ("margin_left","margin_right","margin_top","margin_bottom"):
        setattr(tf,m,Emu(0))
    return tb,tf

def para(tf,text,size,bold,color,italic=False,first=False,align=PP_ALIGN.LEFT,
         space_before=0,space_after=2,line=1.0,bullet=False):
    para=tf.paragraphs[0] if first else tf.add_paragraph()
    para.alignment=align
    para.space_before=Pt(space_before); para.space_after=Pt(space_after)
    try: para.line_spacing=line
    except: pass
    r=para.add_run(); r.text=text
    set_font(r,size,bold,color,italic)
    if bullet:
        pPr=para._pPr if para._pPr is not None else para.get_or_add_pPr()
        # hanging indent for wrapped bullet lines
        pPr.set("indent","-137160"); pPr.set("marL","137160")
    return para

def rect(l,t,w,h,fill,rounded=False,lineclr=None,linew=None):
    shp=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                           Emu(int(l)),Emu(int(t)),Emu(int(w)),Emu(int(h)))
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid(); shp.fill.fore_color.rgb=fill
    if lineclr is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb=lineclr; shp.line.width=Pt(linew or 0.75)
    shp.shadow.inherit=False
    if rounded:
        try: shp.adjustments[0]=0.08
        except: pass
    return shp

# ---------------- geometry ----------------
ML=228600
x0=ML; w_proj=1500000
x1=x0+w_proj; w_del=2620000
x2=x1+w_del; w_tr=2380000
x3=x2+w_tr; w_nf=2186800
cols=[(x0,w_proj),(x1,w_del),(x2,w_tr),(x3,w_nf)]
PAD=64000

# ---------------- title ----------------
_,tf=textbox(ML,155000,7900000,300000)
para(tf,"Product & Tech Updates  —  Carrier Matching",15,True,NAVY,first=True,space_after=1)
_,tf=textbox(ML,455000,8300000,230000)
para(tf,"Apr–Jun'26  ·  Q2 Board Review + Monthly TML Review   |   Steady momentum; pedal on inventory quality & demand fulfilment",
     10,False,GREY,first=True,italic=True)

# ---------------- header row ----------------
hy=690000; hh=300000
hdr=[("PROJECT / GOAL",NAVY),("APR–JUN DELIVERY",GREEN),
     ("TREND  (Apr–Jun'26)",SLATE),("NEXT FOCUS  (Jul–Aug)",PURPLE)]
for (l,w),(txt,clr) in zip(cols,hdr):
    rect(l,hy,w-24000,hh,clr,rounded=True)
    _,tf=textbox(l+40000,hy,w-24000-40000,hh,anchor=MSO_ANCHOR.MIDDLE)
    para(tf,txt,9.5,True,WHITE,first=True,align=PP_ALIGN.LEFT,space_after=0)

# ---------------- rows ----------------
rows=[
 dict(name="Drive FO App to 2,000 DAU", goal="Better load experience & payment visibility",
      chart="row1_app",
      delivery=[("Shipped better load experience & payment visibility",False),
                ("DAU 178 → 252 (+42%); crossed 250/day",True),
                ("MAU 1,027 → 1,325 (+29%); Installs → 1,935",False),
                ("Stickiness 17.3% → 19.0%",False)],
      focus=[("Drive toward 2,000 DAU",False),
             ("Lift WAU → DAU daily habit & retention",False),
             ("Deepen payment visibility + load discovery",False)]),
 dict(name="Improve demand fulfilment to 40%", goal="Matching, bid ranges & real-time visibility",
      chart="row2_demand",
      delivery=[("Matching, bid ranges & live visibility shipped",False),
                ("June rebound — Bids 477 → 725 (+52%)",True),
                ("Placeable fulfilment 11.2% → 17.1%",True),
                ("Unique FO bidders 189 → 267",False)],
      focus=[("Lift placement on inventory & FO bids",False),
             ("Convert bids → trips (Jun dipped to 10)",False),
             ("Push placeable fulfilment toward 40%",False)]),
 dict(name="Increase inventory 50% (1,500/mo)", goal="Via automated agentic calling",
      chart="row3_inventory",
      delivery=[("AI drives ~65% of inventory (869 of 1,341 in Jun)",True),
                ("Agent inv. 728 → 1,043 peak; manual 1,098 → 472",False),
                ("Gap: total inventory & conversion fell 23.9% → 11.4%",False)],
      focus=[("Inventory quality — best-time-to-call, power-lane proximity, FO filtering",False),
             ("Rebuild placeable conversion",False),
             ("Scale agent inventory to 1,500/mo",False)]),
 dict(name="Cost to serve reduction by 70%", goal="Automate ops processes",
      chart="row4_cost",
      delivery=[("Automated Vahan verification & journey updates in LSP WhatsApp [May]",False),
                ("Agent calls scaled to 16k/mo; ~10 hrs/day CT effort saved",True),
                ("~15% reduction in City Team resource cost",True)],
      focus=[("FO document collection via app payment incentives",False),
             ("AI CT for SIM-consent automation",False),
             ("Transit-delay detection",False)]),
 dict(name="Ledger automation for compliance", goal="Trip & party-level ledger",
      chart="row5_ledger",
      delivery=[("Initiative picked up in June",True),
                ("Scoping trip-level & party-level ledger design",False)],
      focus=[("Build trip-level & party-level ledgers",False),
             ("Wallet-based auto-reconciliation across trips",False)]),
]

ry=1030000; rh=738000; gap=16000
for i,row in enumerate(rows):
    top=ry+i*(rh+gap)
    band=BAND if i%2==0 else WHITE
    rect(ML, top, x3+w_nf-ML, rh, band, rounded=True)
    # thin column separators
    # project
    _,tf=textbox(x0+PAD, top+40000, w_proj-2*PAD, rh-70000, anchor=MSO_ANCHOR.MIDDLE)
    para(tf,row["name"],9.5,True,NAVY,first=True,space_after=3,line=1.02)
    para(tf,row["goal"],7.8,False,GREY,italic=True,space_after=0,line=1.0)
    # delivery
    _,tf=textbox(x1+PAD, top+34000, w_del-2*PAD, rh-60000, anchor=MSO_ANCHOR.MIDDLE)
    for j,(txt,hl) in enumerate(row["delivery"]):
        para(tf,"•  "+txt, 8.2, hl, GREEN if hl else INK, first=(j==0),
             space_after=2.5, line=1.02, bullet=True)
    # trend chart
    cpath=f"charts/{row['chart']}.png"
    im=Image.open(cpath); ar=im.size[0]/im.size[1]
    maxW=w_tr-2*PAD; maxH=rh-90000
    cw=maxW; ch=cw/ar
    if ch>maxH: ch=maxH; cw=ch*ar
    cx=x2+(w_tr-cw)/2; cy=top+(rh-ch)/2
    s.shapes.add_picture(cpath, Emu(int(cx)), Emu(int(cy)), Emu(int(cw)), Emu(int(ch)))
    # next focus
    _,tf=textbox(x3+PAD, top+34000, w_nf-2*PAD, rh-60000, anchor=MSO_ANCHOR.MIDDLE)
    for j,(txt,hl) in enumerate(row["focus"]):
        para(tf,"›  "+txt, 8.2, False, PURPLE if j==0 else INK, first=(j==0),
             space_after=2.5, line=1.02, bullet=True)

p.save("deck_out.pptx")
print("saved deck_out.pptx")
